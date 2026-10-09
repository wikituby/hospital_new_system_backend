#!/usr/bin/env python3
from __future__ import annotations
import copy
from pathlib import Path
from docx import Document

def main() -> None:
    src = Path(r"c:\Users\hp\Desktop\hospital documents\Admin documents\compassion invoices\MAY INVOICE KATOMA CDC.docx")
    out = Path(__file__).resolve().parent.parent / "src/main/resources/compassion-invoice/compassion-group-invoice-template.docx"
    doc = Document(str(src))
    body = doc.element.body
    tables = list(doc.tables)
    first_tbl = tables[0]
    for tbl in tables[1:-1]:
        body.remove(tbl._tbl)
    while len(first_tbl.rows) > 4:
        first_tbl._tbl.remove(first_tbl.rows[-1]._tr)
    while len(first_tbl.rows) < 4:
        first_tbl._tbl.append(copy.deepcopy(first_tbl.rows[1]._tr))
    keep = set()
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if i <= 2 or "Prepared By" in text or "Signature" in text:
            keep.add(para._element)
    for para in list(doc.paragraphs):
        if para._element not in keep:
            body.remove(para._element)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print("Saved", out, len(doc.tables), "tables", len(doc.paragraphs), "paras")

if __name__ == "__main__":
    main()
